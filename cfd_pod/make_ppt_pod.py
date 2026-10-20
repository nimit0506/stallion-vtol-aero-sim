"""
Results deck for the payload-pod simulation only (stallion_v2_equipped_with_pod.glb).

  ..\\.venv\\Scripts\\python.exe make_ppt_pod.py     -> ..\\presentations\\Stallion_pod_CFD.pptx

Everything on the slides comes from this case's files: geometry_report.json, config.json,
logs/mesh_log.txt, results/forces.csv, results/aoa_*_forces.out, the Fluent transcript and renders/.
"""
import csv
import json
import re
from pathlib import Path

import numpy as np
from pptx import Presentation
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_MARKER_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
OUT = HERE.parent / "presentations" / "Stallion_pod_CFD.pptx"
RES, REN = HERE / "results", HERE / "renders"

BG = RGBColor(0x1E, 0x20, 0x23)       # matches the Fluent dark-theme renders
PANEL = RGBColor(0x2A, 0x2D, 0x31)
HEAD = RGBColor(0x3A, 0x3E, 0x43)
GRID = RGBColor(0x3A, 0x3E, 0x43)
AXIS = RGBColor(0x55, 0x5A, 0x60)
TEXT = RGBColor(0xF2, 0xF2, 0xF2)
MUTED = RGBColor(0xA8, 0xAE, 0xB5)
ACC = RGBColor(0xFF, 0x8A, 0x3D)      # accent
ANG_COL = {0: RGBColor(0x4C, 0xC9, 0xF0), 4: RGBColor(0xFF, 0x8A, 0x3D), 8: RGBColor(0xE5, 0x4B, 0x6A)}
RES_COL = [RGBColor(0xF2, 0xF2, 0xF2), RGBColor(0x4C, 0xC9, 0xF0), RGBColor(0x2E, 0xC4, 0xB6),
           RGBColor(0x9B, 0x8C, 0xF2), RGBColor(0xFF, 0x8A, 0x3D), RGBColor(0xE5, 0x4B, 0x6A)]
FONT = "Calibri"
W, H = Inches(13.333), Inches(7.5)

# ---------------------------------------------------------------- data
cfg = json.loads((HERE / "config.json").read_text(encoding="utf-8-sig"))
geo = json.loads((HERE / "geometry_report.json").read_text())
V = cfg["flight"]["speed_m_s"]
SOLVER = cfg["solver"]
with open(RES / "forces.csv", newline="") as fh:
    F = {float(r["aoa_deg"]): r for r in csv.DictReader(fh)}
ANGLES = sorted(F)
# cell count as printed by Fluent Meshing at the end of the mesh log
cells = int(re.findall(r"Total cells : (\d+)", (HERE / "logs" / "mesh_log.txt").read_text(errors="ignore"))[-1])


def history(a):
    """Force history of one angle from results/aoa_<a>_forces.out: rows of (iteration, CL, CD, Cm)."""
    rows = []
    for line in (RES / f"aoa_{a:g}_forces.out").read_text().splitlines():
        p = line.split()
        if len(p) == 4 and p[0].isdigit():
            rows.append([float(x) for x in p])
    return np.array(rows)          # iteration (cumulative), cl, cd, cm


HIST = {a: history(a) for a in ANGLES}
# scaled residuals: Fluent prints one line per iteration in its transcript (.trn):
#   iter  continuity  x-vel  y-vel  z-vel  k  omega  ...  -> parse the 7 leading numbers
trn = sorted(RES.glob("fluent-*.trn"))[-1].read_text(errors="ignore")
resid = np.array([[float(x) for x in m.groups()] for m in re.finditer(
    r"^\s*(\d+)\s+([\d.]+e[-+]\d+)\s+([\d.]+e[-+]\d+)\s+([\d.]+e[-+]\d+)\s+([\d.]+e[-+]\d+)"
    r"\s+([\d.]+e[-+]\d+)\s+([\d.]+e[-+]\d+)", trn, re.M)])
RES_NAMES = ["continuity", "x-velocity", "y-velocity", "z-velocity", "k", "omega"]


def g(a, k):
    """One value from forces.csv as float."""
    return float(F[a][k])


def spread_last(a, col, n=50):
    """(max - min) / |mean| of the last n iterations, in %"""
    x = HIST[a][-n:, col]
    return (x.max() - x.min()) / abs(x.mean()) * 100


def render(a, scene, view):
    """Path of a Fluent still from 04_demo_setup.py; fail early if it was never rendered."""
    p = REN / f"aoa{a:g}_{scene}_{view}.png"
    if not p.exists():
        raise FileNotFoundError(p)
    return p


# ---------------------------------------------------------------- helpers
def new_slide(prs, notes=None):
    """Blank dark slide; notes go into the speaker notes."""
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = BG
    if notes:
        s.notes_slide.notes_text_frame.text = notes
    return s


def text(slide, x, y, w, h, runs, size=16, color=TEXT, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    """Text box. runs: str, or list of paragraphs; a paragraph is a str or a list of (text, {size, bold, italic, color})."""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, para in enumerate(runs if isinstance(runs, list) else [runs]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(6)
        for t, o in ([(para, {})] if isinstance(para, str) else para):
            r = p.add_run()
            r.text = t
            f = r.font
            f.name, f.size = FONT, Pt(o.get("size", size))
            f.bold, f.italic = o.get("bold", bold), o.get("italic", False)
            f.color.rgb = o.get("color", color)
    return tb


def title(slide, t, sub=None):
    """Slide title (and optional grey subtitle) in the fixed top-left position."""
    text(slide, Inches(0.6), Inches(0.4), Inches(12.1), Inches(0.7), t, size=34, bold=True)
    if sub:
        text(slide, Inches(0.6), Inches(1.1), Inches(12.1), Inches(0.4), sub, size=16, color=MUTED)


def image(slide, path, x, y, w=None, h=None, caption=None):
    """Picture (give w or h to keep the aspect ratio) with an optional centred caption below it."""
    pic = slide.shapes.add_picture(str(path), x, y, w, h)
    if caption:
        text(slide, x, pic.top + pic.height + Inches(0.06), pic.width, Inches(0.3), caption,
             size=12, color=MUTED, align=PP_ALIGN.CENTER)
    return pic


def panel(slide, x, y, w, h):
    """Flat rounded card behind text (no outline, no shadow)."""
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    sh.adjustments[0] = 0.06
    sh.fill.solid()
    sh.fill.fore_color.rgb = PANEL
    sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def table(slide, rows, x, y, widths, row_h=Inches(0.42), size=13, color_fn=None):
    """Dark-theme table; row 0 is the header. color_fn(row, col, value) can recolour single cells."""
    t = slide.shapes.add_table(len(rows), len(rows[0]), x, y, sum(widths), row_h * len(rows)).table
    for c, wd in enumerate(widths):
        t.columns[c].width = wd
    for r, row in enumerate(rows):
        t.rows[r].height = row_h
        for c, val in enumerate(row):
            cell = t.cell(r, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = HEAD if r == 0 else PANEL
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.text = str(val) or " "
            f = cell.text_frame.paragraphs[0].runs[0].font
            f.name, f.size, f.bold = FONT, Pt(size), r == 0
            f.color.rgb = (color_fn(r, c, val) if color_fn and r else None) or TEXT
    return t


def iteration_axis(ch):
    """x axis 0..last iteration, labels below the plot (not where y crosses 1 on a log axis)"""
    from pptx.enum.chart import XL_TICK_LABEL_POSITION
    ch.category_axis.minimum_scale = 0
    ch.category_axis.maximum_scale = float(resid[-1, 0])
    ch.category_axis.major_unit = 150
    ch.category_axis.tick_label_position = XL_TICK_LABEL_POSITION.LOW


def style_axes(ch, size=11):
    """Muted tick labels and axis lines, horizontal gridlines only."""
    for ax in (ch.category_axis, ch.value_axis):
        ax.tick_labels.font.size = Pt(size)
        ax.tick_labels.font.color.rgb = MUTED
        ax.tick_labels.font.name = FONT
        ax.format.line.color.rgb = AXIS
    ch.value_axis.has_major_gridlines = True
    ch.value_axis.major_gridlines.format.line.color.rgb = GRID


def chart_title(ch, t, size=15):
    """Bold white chart title."""
    ch.has_title = True
    ch.chart_title.text_frame.text = t
    f = ch.chart_title.text_frame.paragraphs[0].runs[0].font
    f.size, f.bold, f.color.rgb, f.name = Pt(size), True, TEXT, FONT


def legend(ch):
    """Legend under the plot area, outside the plot."""
    ch.has_legend = True
    ch.legend.position = XL_LEGEND_POSITION.BOTTOM
    ch.legend.include_in_layout = False
    ch.legend.font.size, ch.legend.font.color.rgb, ch.legend.font.name = Pt(11), MUTED, FONT


def axis_title(ax, t):
    """Small grey axis title."""
    ax.has_title = True
    ax.axis_title.text_frame.text = t
    f = ax.axis_title.text_frame.paragraphs[0].runs[0].font
    f.size, f.color.rgb, f.name, f.bold = Pt(11), MUTED, FONT, False


def polar_chart(slide, x, y, w, h, ttl, key, fmt):
    """Native line chart of one forces.csv column against angle of attack, values labelled."""
    cd = CategoryChartData()
    cd.categories = [f"{a:g}°" for a in ANGLES]
    cd.add_series(ttl, [g(a, key) for a in ANGLES])
    ch = slide.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, x, y, w, h, cd).chart
    chart_title(ch, ttl)
    ch.has_legend = False
    style_axes(ch)
    ch.value_axis.tick_labels.number_format = fmt
    ch.value_axis.tick_labels.number_format_is_linked = False
    s = ch.series[0]
    s.smooth = False
    s.format.line.color.rgb, s.format.line.width = ACC, Pt(2.5)
    s.marker.style, s.marker.size = XL_MARKER_STYLE.CIRCLE, 8
    s.marker.format.fill.solid()
    s.marker.format.fill.fore_color.rgb = ACC
    s.marker.format.line.color.rgb = ACC
    dl = s.data_labels
    dl.show_value, dl.number_format, dl.number_format_is_linked = True, fmt, False
    dl.position = XL_LABEL_POSITION.ABOVE
    dl.font.size, dl.font.color.rgb, dl.font.name = Pt(12), TEXT, FONT
    return ch


def history_chart(slide, x, y, w, h, ttl, col, fmt):
    """XY chart of one force coefficient against cumulative iteration, one line per angle (col 1 = CL, 2 = CD)."""
    cd = XyChartData()
    for a in ANGLES:
        s = cd.add_series(f"{a:g}°")
        for row in HIST[a]:
            s.add_data_point(row[0], row[col])
    ch = slide.shapes.add_chart(XL_CHART_TYPE.XY_SCATTER_LINES_NO_MARKERS, x, y, w, h, cd).chart
    chart_title(ch, ttl)
    legend(ch)
    style_axes(ch)
    ch.value_axis.tick_labels.number_format = fmt
    ch.value_axis.tick_labels.number_format_is_linked = False
    ch.category_axis.has_major_gridlines = False
    axis_title(ch.category_axis, "iteration (cumulative; each angle starts from the previous solution)")
    iteration_axis(ch)
    for s, a in zip(ch.series, ANGLES):
        s.smooth = False
        s.format.line.color.rgb, s.format.line.width = ANG_COL[int(a)], Pt(1.75)
        s.marker.style = XL_MARKER_STYLE.NONE
    return ch


def residual_chart(slide, x, y, w, h):
    """Fluent's six scaled residuals against iteration, log10 y axis (python-pptx has no log option, so it is set in the XML)."""
    cd = XyChartData()
    step = 3                                   # thin the 900 points a little; shape unchanged
    for j, name in enumerate(RES_NAMES):
        s = cd.add_series(name)
        for row in resid[::step]:
            s.add_data_point(row[0], row[j + 1])
    ch = slide.shapes.add_chart(XL_CHART_TYPE.XY_SCATTER_LINES_NO_MARKERS, x, y, w, h, cd).chart
    chart_title(ch, "Scaled residuals (Fluent)")
    legend(ch)
    style_axes(ch)
    ch.category_axis.has_major_gridlines = False
    ch.value_axis.tick_labels.number_format = "0.E+0"
    ch.value_axis.tick_labels.number_format_is_linked = False
    scaling = ch.value_axis._element.find(qn("c:scaling"))      # log10 y axis
    lb = OxmlElement("c:logBase")
    lb.set("val", "10")
    scaling.insert(0, lb)
    axis_title(ch.category_axis, "iteration (0° 1–150 · 4° 151–300 · 8° 301–900)")
    iteration_axis(ch)
    for s, col in zip(ch.series, RES_COL):
        s.smooth = False
        s.format.line.color.rgb, s.format.line.width = col, Pt(1.25)
        s.marker.style = XL_MARKER_STYLE.NONE
    return ch


# ---------------------------------------------------------------- deck
prs = Presentation()
prs.slide_width, prs.slide_height = W, H
GAP = Inches(0.15)

# 1 title -------------------------------------------------------------
s = new_slide(prs, "CFD of the Stallion VTOL equipped with the payload pod, built from stallion_v2_equipped_with_pod.glb.")
hero = image(s, render(4, "2-surface-pressure", "iso"), 0, 0, h=Inches(5.3))
hero.crop_left = 0.27                                  # crop the colour bar off the left of the render
hero.width = int(hero.height * (16 / 9) * (1 - 0.27))  # keep the aspect ratio after cropping
hero.left = W - hero.width                             # flush with the right edge
text(s, Inches(0.6), Inches(2.0), Inches(5.6), Inches(1.6),
     [[("Stallion VTOL · payload pod", {"size": 22, "color": MUTED})],
      [("External aero CFD", {"size": 44, "bold": True})]])
text(s, Inches(0.6), Inches(5.6), Inches(12), Inches(1.2),
     [f"{V:g} m/s · angle of attack {', '.join(f'{a:g}' for a in ANGLES)}°",
      [("Ansys Fluent 2026 R1 Student · steady RANS, k-ω SST", {"color": MUTED, "size": 14})]], size=18)

# 2 results table -------------------------------------------------------
s = new_slide(prs, "Forces integrated over the whole aircraft surface. 8° did not meet the convergence criterion "
                   "within 600 iterations; its row is the mean of the last 200 iterations.")
title(s, "Results", f"Whole aircraft at {V:g} m/s, sea-level air (ρ {cfg['flight']['density_kg_m3']} kg/m³)")
rows = [("AoA", "CL", "CD", "Cm", "L/D", "Lift (N)", "Drag (N)", "Iterations", "Status")]
for a in ANGLES:
    r = F[a]
    rows.append((f"{a:g}°", f"{g(a,'CL'):.3f}", f"{g(a,'CD'):.4f}", f"{g(a,'Cm') + 0.0:.3f}".replace("-0.000", "0.000"),
                 f"{g(a,'L_over_D'):.2f}", f"{g(a,'lift_N'):.1f}", f"{g(a,'drag_N'):.2f}", r["iterations"],
                 "converged" if r["status"] == "converged" else "not converged (mean of last 200 it.)"))
table(s, rows, Inches(0.6), Inches(1.85), [Inches(w) for w in (0.85, 1.05, 1.15, 1.05, 0.95, 1.15, 1.15, 1.2, 3.55)],
      row_h=Inches(0.55), size=15,
      color_fn=lambda r, c, v: ACC if (c == 8 and "not" in v) else None)
ref = cfg["reference"]
for i, (val, lab) in enumerate([
        (f"{max(g(a,'L_over_D') for a in ANGLES):.1f}", f"best L/D (at {max(ANGLES, key=lambda a: g(a,'L_over_D')):g}°)"),
        (f"{g(4,'lift_N'):.1f} N", "lift at 4°"),
        (f"{g(4,'drag_N'):.2f} N", "drag at 4°")]):
    x = Inches(0.6) + i * Inches(4.1)
    panel(s, x, Inches(4.35), Inches(3.9), Inches(1.6))
    text(s, x + Inches(0.25), Inches(4.5), Inches(3.4), Inches(0.8), val, size=40, bold=True, color=ACC)
    text(s, x + Inches(0.25), Inches(5.35), Inches(3.4), Inches(0.5), lab, size=14, color=MUTED)
text(s, Inches(0.6), Inches(6.35), Inches(12.1), Inches(0.6),
     f"Reference area {ref['area_m2']} m², reference length {ref['length_m']} m, moment centre "
     f"{ref['moment_center_m']} m (wing quarter chord). Cm positive nose-up.", size=12, color=MUTED)

# 3 coefficient charts ----------------------------------------------------
s = new_slide(prs, "Lift, drag and lift-to-drag ratio against angle of attack.")
title(s, "Coefficients vs angle of attack", f"Whole aircraft, {V:g} m/s")
cw = Inches(3.95)
for i, (key, ttl, fmt) in enumerate([("CL", "Lift coefficient CL", "0.00"), ("CD", "Drag coefficient CD", "0.000"),
                                     ("L_over_D", "Lift-to-drag L/D", "0.0")]):
    polar_chart(s, Inches(0.6) + i * (cw + Inches(0.13)), Inches(1.75), cw, Inches(4.6), ttl, key, fmt)
text(s, Inches(0.6), Inches(6.55), Inches(12.1), Inches(0.5),
     "8° value is the mean of the last 200 iterations (not converged within 600).", size=12, color=MUTED)

# 4 geometry ------------------------------------------------------------
raw, wr = geo["raw"], geo["wrapped"]
s = new_slide(prs, "The GLB could not be meshed as delivered, so it was repaired and shrink-wrapped into one "
                   "watertight outer skin before meshing.")
title(s, "Geometry preparation", "stallion_v2_equipped_with_pod.glb → watertight skin")
rows = [("", "Input GLB", "Wrapped skin"),
        ("CAD parts", "96", "1 body"),
        ("Separate bodies", f"{geo['bodies']}", "1"),
        ("Non-manifold edges", f"{raw['non_manifold_edges']:,}", f"{wr['non_manifold_edges']}"),
        ("Facets", f"{geo['input_facets']:,}", f"{wr['facets']:,}"),
        ("Fuselage interior", "hollow (open hatches)", "solid"),
        ("Max bounding-box change", "", f"{wr['max_bbox_change_mm']:.1f} mm")]
table(s, rows, Inches(0.6), Inches(1.85), [Inches(2.5), Inches(1.95), Inches(1.85)], row_h=Inches(0.5), size=14,
      color_fn=lambda r, c, v: ACC if c == 2 else None)
text(s, Inches(0.6), Inches(5.6), Inches(6.3), Inches(1.5),
     [[("Steps: ", {"bold": True}),
       ("merge the 96 parts and rotate to the solver frame (nose +Y, up +Z) → voxel shrink-wrap at "
        f"{wr['voxel_mm']} mm with openings up to 24 mm sealed → decimate → check: closed, manifold, "
        "and a ray test finds no air inside the fuselage.", {})]], size=13, color=MUTED)
image(s, render(4, "1-geometry-mesh", "mesh-closeup"), Inches(7.15), Inches(1.85), w=Inches(5.6),
      caption="Wrapped skin with the surface mesh, and the centre-plane volume mesh")

# 5 mesh & solver ---------------------------------------------------------
m, fl, dm = cfg["mesh"], cfg["flight"], cfg["domain"]
s = new_slide(prs, "Mesh and solver settings used for all three angles.")
title(s, "Mesh and solver setup")
items = [("Mesh", f"poly-hexcore, {cells:,} cells; {m['boundary_layers']} prism layers, first cell "
                  f"{m['bl_first_height_mm']} mm; surface {m['min_size_mm']:g}–{m['aircraft_max_size_mm']:g} mm"),
         ("Domain", f"box {dm['upstream_m']:g} m upstream, {dm['downstream_m']:g} m downstream, "
                    f"{dm['lateral_m']:g} m sides, {dm['vertical_m']:g} m top/bottom"),
         ("Boundaries", "velocity inlet on 5 faces, pressure outlet behind; aircraft = no-slip wall"),
         ("Flow", f"{V:g} m/s, ρ {fl['density_kg_m3']} kg/m³, μ {fl['viscosity_kg_m_s']} Pa·s, "
                  f"turbulence intensity {fl['turbulence_intensity']*100:g}%"),
         ("Model", "steady, pressure-based coupled solver, k-ω SST"),
         ("Angles", f"{', '.join(f'{a:g}' for a in ANGLES)}°, set by tilting the free-stream direction; "
                    "same mesh for every angle")]
y = Inches(1.6)
for k, v in items:
    text(s, Inches(0.6), y, Inches(1.6), Inches(0.75), k, size=15, bold=True, color=ACC)
    text(s, Inches(2.25), y, Inches(4.45), Inches(0.8), v, size=14)
    y += Inches(0.86)
image(s, render(4, "7-wall-yplus", "top"), Inches(7.0), Inches(1.6), w=Inches(5.75),
      caption="Wall y+ on the aircraft surface (4°)")

# 6 convergence summary -----------------------------------------------------
s = new_slide(prs, "Convergence criterion: CL and CD each change by less than "
                   f"{SOLVER['force_tolerance']*100:g}% over a {SOLVER['check_every']}-iteration block, with "
                   f"{SOLVER['min_iterations']}–{SOLVER['iterations']} iterations per angle.")
title(s, "Convergence by angle",
      f"Criterion: CL and CD change < {SOLVER['force_tolerance']*100:g}% per {SOLVER['check_every']} iterations "
      f"({SOLVER['min_iterations']} min, {SOLVER['iterations']} max)")
rows = [("AoA", "Iterations", "Status", "CL spread, last 50 it.", "CD spread, last 50 it.", "Reported value")]
for a in ANGLES:
    ok = F[a]["status"] == "converged"
    rows.append((f"{a:g}°", F[a]["iterations"], "converged" if ok else "not converged",
                 f"{spread_last(a, 1):.2f}%", f"{spread_last(a, 2):.2f}%",
                 "final iteration" if ok else "mean of last 200 iterations"))
table(s, rows, Inches(0.6), Inches(1.85), [Inches(w) for w in (0.9, 1.35, 1.9, 2.55, 2.55, 2.85)],
      row_h=Inches(0.55), size=15, color_fn=lambda r, c, v: ACC if v == "not converged" else None)
text(s, Inches(0.6), Inches(4.25), Inches(12.1), Inches(2.6),
     ["Spread = (max − min) / mean of the force coefficient over the last 50 iterations.",
      "0° and 4° met the criterion at the minimum iteration count. At 8° CL and CD kept oscillating by several "
      "percent up to the 600-iteration limit; the reported 8° numbers are the average of the last 200 iterations.",
      "Each angle was started from the previous angle's solution (0° → 4° → 8°)."], size=15)

# 7 convergence histories ---------------------------------------------------
s = new_slide(prs, "Force-coefficient histories for every iteration of the run, and Fluent's scaled residuals.")
title(s, "Convergence histories")
cw2 = Inches(5.95)
history_chart(s, Inches(0.6), Inches(1.3), cw2, Inches(2.85), "Lift coefficient CL", 1, "0.00")
history_chart(s, Inches(0.6), Inches(4.25), cw2, Inches(2.85), "Drag coefficient CD", 2, "0.000")
residual_chart(s, Inches(6.8), Inches(1.3), cw2, Inches(5.8))

# 8 surface pressure across angles -----------------------------------------
s = new_slide(prs, "Pressure coefficient on the aircraft surface; same colour scale (−1.5 to 1.0) at every angle. "
                   "Blue = below free-stream pressure, red = above.")
title(s, "Surface pressure at each angle", "Pressure coefficient, same colour scale (−1.5 to 1.0)")
iw = Inches(3.95)
for i, a in enumerate(ANGLES):
    x = Inches(0.6) + i * (iw + Inches(0.13))
    image(s, render(a, "2-surface-pressure", "iso"), x, Inches(1.75), w=iw, caption=f"{a:g}° · top-side view")
    image(s, render(a, "2-surface-pressure", "pod"), x, Inches(4.45), w=iw, caption=f"{a:g}° · pod, from below")

# 9-11 one slide per angle ----------------------------------------------------
for a in ANGLES:
    r = F[a]
    st = "converged" if r["status"] == "converged" else "not converged, mean of last 200 iterations"
    s = new_slide(prs, f"All views at {a:g}° angle of attack. CL {g(a,'CL'):.3f}, CD {g(a,'CD'):.4f}, "
                       f"L/D {g(a,'L_over_D'):.2f} ({st}).")
    title(s, f"Angle of attack {a:g}°",
          f"CL {g(a,'CL'):.3f} · CD {g(a,'CD'):.4f} · L/D {g(a,'L_over_D'):.2f} · lift {g(a,'lift_N'):.1f} N · "
          f"drag {g(a,'drag_N'):.2f} N · {st}")
    shots = [(("2-surface-pressure", "top"), "Pressure coefficient, top"),
             (("2-surface-pressure", "bottom"), "Pressure coefficient, underside"),
             (("3-streamlines", "iso"), "Streamlines, coloured by speed"),
             (("4-wing-section", "section"), "Speed on a wing section at 0.55 m span"),
             (("5-tip-vortices", "tip"), "Tip streamlines, coloured by streamwise vorticity"),
             (("6-wake", "wake"), "Streamwise vorticity, 0.1 m and 0.4 m behind the wing")]
    for k, ((scene, view), cap) in enumerate(shots):
        x = Inches(0.6) + (k % 3) * (iw + Inches(0.13))
        y = Inches(1.65) + (k // 3) * Inches(2.72)
        image(s, render(a, scene, view), x, y, w=iw, caption=cap)

# 12 streamlines around the pod ---------------------------------------------
s = new_slide(prs, "Streamlines released ahead of the nose, viewed from below at the pod.")
title(s, "Flow past the pod", "Streamlines coloured by speed (14–26 m/s), viewed from below")
for i, a in enumerate(ANGLES):
    image(s, render(a, "3-streamlines", "pod"), Inches(0.6) + i * (iw + Inches(0.13)), Inches(1.9), w=iw,
          caption=f"{a:g}°")
image(s, render(4, "2-surface-pressure", "front"), Inches(0.6), Inches(4.55), h=Inches(2.35),
      caption="Pressure coefficient, front view (4°)")
image(s, render(4, "6-wake", "rear"), Inches(4.85), Inches(4.55), h=Inches(2.35),
      caption="Wake planes from behind (4°)")
text(s, Inches(9.25), Inches(4.65), Inches(3.5), Inches(2.3),
     [[("Colour scales", {"bold": True})],
      "Streamlines: 14–26 m/s", "Pressure coefficient: −1.5 to 1.0", "Streamwise vorticity: −40 to 40 1/s"],
     size=13, color=MUTED)

# 13 modelling notes --------------------------------------------------------
s = new_slide(prs, "What was and was not included in this simulation.")
title(s, "Modelling notes")
panel(s, Inches(0.6), Inches(1.6), Inches(12.1), Inches(3.0))
text(s, Inches(0.95), Inches(1.9), Inches(11.4), Inches(3.8),
     [f"Steady RANS (k-ω SST) on a {cells:,}-cell mesh; the Ansys Student licence limits the mesh to about 1M cells.",
      "Airframe, motors and pod as in the GLB; propellers are not modelled.",
      "Openings up to 24 mm (hatch gaps, slots) were sealed by the shrink-wrap; internal parts were removed.",
      "Same mesh for all angles; angle of attack set by the free-stream direction.",
      "8° did not converge within 600 iterations; its forces are the mean of the last 200 iterations."],
     size=16)
text(s, Inches(0.6), Inches(4.9), Inches(12.1), Inches(0.6),
     "Files: cfd_pod/results (cases, data, force histories, forces.csv) · cfd_pod/renders (all images) · "
     "demo.bat opens the case in Fluent", size=12, color=MUTED)

prs.save(OUT)
print(f"wrote {OUT}  ({len(prs.slides)} slides)")
