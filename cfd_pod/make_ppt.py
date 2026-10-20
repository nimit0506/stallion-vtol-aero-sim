"""
Build the results deck: Stallion VTOL with payload pod vs the clean airframe.

  ..\\.venv\\Scripts\\python.exe make_ppt.py        -> ..\\presentations\\Stallion_pod_CFD_results.pptx

Reads results/forces.csv + renders/ of this case and of ..\\cfd_clean (clean airframe).
"""
import csv
import json
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_MARKER_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Emu, Inches, Pt

HERE = Path(__file__).resolve().parent
CLEAN = HERE.parent / "cfd_clean"
OUT = HERE.parent / "presentations" / "Stallion_pod_CFD_results.pptx"

BG = RGBColor(0x1E, 0x20, 0x23)       # matches the Fluent dark-theme renders
PANEL = RGBColor(0x2A, 0x2D, 0x31)
TEXT = RGBColor(0xF2, 0xF2, 0xF2)
MUTED = RGBColor(0xA8, 0xAE, 0xB5)
POD = RGBColor(0xFF, 0x8A, 0x3D)      # pod case = warm accent
CLN = RGBColor(0x4C, 0xC9, 0xF0)      # clean case = cool
FONT = "Calibri"
W, H = Inches(13.333), Inches(7.5)


def forces(case_dir):
    """forces.csv of a case as {aoa: row dict}."""
    with open(case_dir / "results" / "forces.csv", newline="") as fh:
        return {float(r["aoa_deg"]): r for r in csv.DictReader(fh)}


def render(case_dir, aoa, scene, view):
    """Path of a Fluent still from 04_demo_setup.py; fail early if it was never rendered."""
    p = case_dir / "renders" / f"aoa{aoa:g}_{scene}_{view}.png"
    if not p.exists():
        raise FileNotFoundError(p)
    return p


# ---------------------------------------------------------------- helpers
def new_slide(prs, notes=None):
    """Blank dark slide; notes go into the speaker notes."""
    s = prs.slides.add_slide(prs.slide_layouts[6])          # blank
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = BG
    if notes:
        s.notes_slide.notes_text_frame.text = notes
    return s


def text(slide, x, y, w, h, runs, size=16, color=TEXT, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, name=None):
    """runs: str, or list of paragraphs; a paragraph is str or list of (text, {opts})"""
    tb = slide.shapes.add_textbox(x, y, w, h)
    if name:
        tb.name = name                                        # named shapes are easier to find in PowerPoint
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0   # align text with images
    tf.vertical_anchor = anchor
    paras = runs if isinstance(runs, list) else [runs]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(6)
        for t, o in ([(para, {})] if isinstance(para, str) else para):
            r = p.add_run()
            r.text = t
            f = r.font
            f.name = FONT
            f.size = Pt(o.get("size", size))
            f.bold = o.get("bold", bold)
            f.italic = o.get("italic", False)
            f.color.rgb = o.get("color", color)
    return tb


def title(slide, t, sub=None):
    """Slide title (and optional grey subtitle) in the fixed top-left position."""
    text(slide, Inches(0.6), Inches(0.4), Inches(12.1), Inches(0.7), t, size=34, bold=True, name="Title")
    if sub:
        text(slide, Inches(0.6), Inches(1.1), Inches(12.1), Inches(0.4), sub, size=16, color=MUTED, name="Subtitle")


def image(slide, path, x, y, w=None, h=None, caption=None):
    """Picture (give w or h to keep its aspect ratio) with an optional centred caption below it."""
    pic = slide.shapes.add_picture(str(path), x, y, w, h)
    if caption:
        text(slide, x, pic.top + pic.height + Inches(0.08), pic.width, Inches(0.35), caption,
             size=12, color=MUTED, align=PP_ALIGN.CENTER)
    return pic


def panel(slide, x, y, w, h):
    """Flat rounded card behind text (no outline, no shadow)."""
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    sh.adjustments[0] = 0.06                                  # corner radius
    sh.fill.solid()
    sh.fill.fore_color.rgb = PANEL
    sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def stat(slide, x, y, w, value, label, color):
    """Big-number card: value on top, explanation underneath."""
    panel(slide, x, y, w, Inches(1.75))
    text(slide, x + Inches(0.25), y + Inches(0.18), w - Inches(0.5), Inches(0.9), value,
         size=44, bold=True, color=color)
    text(slide, x + Inches(0.25), y + Inches(1.05), w - Inches(0.5), Inches(0.6), label,
         size=14, color=MUTED)


def line_chart(slide, x, y, w, h, title_txt, series, fmt="0.00", label_pos=None):
    """Native (editable) PowerPoint line chart over the angles; series = [(name, values)], clean then pod.
    label_pos puts each series' value labels above/below so the two series' labels don't collide."""
    from pptx.enum.chart import XL_LABEL_POSITION as LP
    label_pos = label_pos or (LP.ABOVE, LP.BELOW)
    cd = CategoryChartData()
    cd.categories = [f"{a:g}°" for a in ANGLES]
    for name, vals in series:
        cd.add_series(name, vals)
    gf = slide.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, x, y, w, h, cd)
    ch = gf.chart
    ch.has_title = True
    ch.chart_title.text_frame.text = title_txt
    tp = ch.chart_title.text_frame.paragraphs[0].runs[0].font
    tp.size, tp.bold, tp.color.rgb, tp.name = Pt(16), True, TEXT, FONT
    # dark-theme styling: muted axes and grid, legend under the plot
    ch.has_legend = True
    ch.legend.position = XL_LEGEND_POSITION.BOTTOM
    ch.legend.include_in_layout = False
    ch.legend.font.size, ch.legend.font.color.rgb, ch.legend.font.name = Pt(12), MUTED, FONT
    for ax in (ch.category_axis, ch.value_axis):
        ax.tick_labels.font.size = Pt(12)
        ax.tick_labels.font.color.rgb = MUTED
        ax.tick_labels.font.name = FONT
        ax.format.line.color.rgb = RGBColor(0x55, 0x5A, 0x60)
    ch.value_axis.has_major_gridlines = True
    ch.value_axis.major_gridlines.format.line.color.rgb = RGBColor(0x3A, 0x3E, 0x43)
    ch.value_axis.tick_labels.number_format = fmt
    ch.value_axis.tick_labels.number_format_is_linked = False
    # per-series colour, markers and value labels (clean = cool blue, pod = orange)
    for s, col, pos in zip(ch.series, (CLN, POD), (label_pos[0], label_pos[1])):
        s.format.line.color.rgb = col
        s.format.line.width = Pt(2.5)
        s.smooth = False
        s.marker.style = XL_MARKER_STYLE.CIRCLE
        s.marker.size = 8
        s.marker.format.fill.solid()
        s.marker.format.fill.fore_color.rgb = col
        s.marker.format.line.color.rgb = col
        s.data_labels.show_value = True
        s.data_labels.number_format = fmt
        s.data_labels.number_format_is_linked = False
        s.data_labels.font.size = Pt(11)
        s.data_labels.font.color.rgb = col
        s.data_labels.font.name = FONT
        s.data_labels.position = pos
    return ch


def pct(a, b):
    """Percentage change from b to a."""
    return (a - b) / b * 100.0


# ---------------------------------------------------------------- deck
# inputs: both cases' forces, the wrap report, the cell count from the mesher's log, the flight speed
fp, fc = forces(HERE), forces(CLEAN)
ANGLES = sorted(fp)
geo = json.loads((HERE / "geometry_report.json").read_text())
import re
_cells = re.findall(r"Total cells : (\d+)", (HERE / "logs" / "mesh_log.txt").read_text(errors="ignore"))
mesh_rep = {"cells": int(_cells[-1])} if _cells else {}
cfg = json.loads((HERE / "config.json").read_text(encoding="utf-8-sig"))
V = cfg["flight"]["speed_m_s"]


def g(tbl, a, k):
    """One value from a forces table as float."""
    return float(tbl[a][k])


prs = Presentation()
prs.slide_width, prs.slide_height = W, H                  # 16:9

# 1 title
s = new_slide(prs, "Payload-pod variant of the Stallion VTOL, same CFD pipeline and settings as the clean airframe, "
                   "so the differences are the pod and fitted motors.")
hero = image(s, render(HERE, 4, "2-surface-pressure", "iso"), Inches(5.9), Inches(0.0), h=Inches(5.3))
hero.crop_left = 0.27                              # drop the legend; keep the aircraft and logo
hero.width = int(hero.height * (16 / 9) * (1 - 0.27))
hero.left = W - hero.width                         # flush with the right edge
text(s, Inches(0.6), Inches(2.0), Inches(5.4), Inches(1.6),
     [[("Stallion VTOL", {"size": 24, "color": MUTED})],
      [("Payload-pod CFD", {"size": 44, "bold": True})]], size=44)
text(s, Inches(0.6), Inches(5.6), Inches(12), Inches(1.2),
     [f"External aerodynamics at {V:g} m/s, angle of attack {', '.join(f'{a:g}' for a in ANGLES)}°",
      [("Ansys Fluent 2026 R1 Student · k-ω SST · compared with the clean airframe", {"color": MUTED, "size": 14})]],
     size=18)

# 2 key results
a = 4.0
s = new_slide(prs, "Headline numbers at the cruise angle (4°). Coefficients use the same reference area, chord "
                   "and moment centre for both configurations.")
title(s, "The pod costs drag, not lift", f"Cruise point: {a:g}° angle of attack, {V:g} m/s")
dcd = pct(g(fp, a, "CD"), g(fc, a, "CD"))
dld = pct(g(fp, a, "L_over_D"), g(fc, a, "L_over_D"))
dcl = pct(g(fp, a, "CL"), g(fc, a, "CL"))
ddrag = g(fp, a, "drag_N") - g(fc, a, "drag_N")
cw = Inches(2.85)
for i, (val, lab, col) in enumerate([
        (f"{dcd:+.0f}%", f"drag coefficient\n{g(fc,a,'CD'):.4f} → {g(fp,a,'CD'):.4f}", POD),
        (f"{ddrag:+.2f} N", f"extra drag at {V:g} m/s\n{g(fc,a,'drag_N'):.2f} → {g(fp,a,'drag_N'):.2f} N", POD),
        (f"{dld:+.0f}%", f"lift-to-drag ratio\n{g(fc,a,'L_over_D'):.1f} → {g(fp,a,'L_over_D'):.1f}", POD),
        (f"{dcl:+.0f}%", f"lift coefficient\n{g(fc,a,'CL'):.3f} → {g(fp,a,'CL'):.3f}", CLN)]):
    stat(s, Inches(0.6) + i * (cw + Inches(0.3)), Inches(1.9), cw, val, lab, col)
image(s, render(HERE, a, "2-surface-pressure", "pod"), Inches(0.6), Inches(4.0), h=Inches(3.0))
text(s, Inches(6.2), Inches(4.2), Inches(6.5), Inches(2.8),
     [[("What the numbers say", {"bold": True, "size": 18})],
      "The blunt pod (and the boom-front motors) add frontal area and wetted surface, so drag rises at every angle.",
      "Lift is nearly unchanged: the wing does the lifting and the pod barely alters its loading.",
      [("Clean = original airframe STL · Pod = equipped GLB: payload pod plus tilt motors on the booms",
        {"size": 12, "color": MUTED})]],
     size=15)

# 3 geometry
raw = geo["raw"]
s = new_slide(prs, "The GLB is a CAD assembly: 96 parts including internal electronics and spars, many touching or "
                   "overlapping. It cannot be meshed directly, so it was shrink-wrapped into one solid skin.")
title(s, "Geometry: 96-part CAD model → one watertight skin", "Repair test on stallion_v2_equipped_with_pod.glb")
rows = [("", "Raw GLB", "Wrapped skin"),
        ("Separate parts / bodies", "96 / 251", "1"),
        ("Non-manifold edges", f"{raw['non_manifold_edges']:,}", "0"),
        ("Open edges", "76 (in GLB)", "0"),
        ("Fuselage interior", "hollow, open hatches", "solid (gaps ≤ 24 mm sealed)"),
        ("Facets", f"{geo['input_facets']:,}", f"{geo['wrapped']['facets']:,}"),
        ("Max shape change", "", f"{geo['wrapped']['max_bbox_change_mm']:.1f} mm")]
tbl = s.shapes.add_table(len(rows), 3, Inches(0.6), Inches(1.8), Inches(6.4), Inches(3.6)).table
tbl.columns[0].width, tbl.columns[1].width, tbl.columns[2].width = Inches(2.4), Inches(1.9), Inches(2.1)
for r, row in enumerate(rows):
    for c, val in enumerate(row):
        cell = tbl.cell(r, c)
        cell.fill.solid()
        cell.fill.fore_color.rgb = PANEL if r else RGBColor(0x3A, 0x3E, 0x43)
        cell.text = val
        para = cell.text_frame.paragraphs[0]
        f = para.runs[0].font if para.runs else para.font
        f.size, f.name, f.bold = Pt(13), FONT, r == 0 or c == 0
        f.color.rgb = (POD if c == 2 and r else TEXT)
text(s, Inches(0.6), Inches(5.65), Inches(6.4), Inches(1.4),
     "Voxel shrink-wrap at 1.5 mm: air flood-filled from outside, signed-distance skin from the real facets, "
     "then decimated to the Student-licence size. A ray test confirms no air pockets inside the body.",
     size=13, color=MUTED)
image(s, render(HERE, 4, "1-geometry-mesh", "mesh-closeup"), Inches(7.3), Inches(1.8), w=Inches(5.45),
      caption="Wrapped skin and the CFD surface mesh")

# 4 mesh & setup
m = cfg["mesh"]
s = new_slide(prs, "Identical mesh sizing and solver settings for both configurations.")
title(s, "Mesh and solver setup", "Same settings for both configurations; only the geometry changes")
image(s, render(HERE, 4, "7-wall-yplus", "top"), Inches(6.9), Inches(1.8), w=Inches(5.85),
      caption="Wall y+ ≈ 5–25: SST with wall treatment")
_cc = re.findall(r"Total cells : (\d+)", (CLEAN / "logs" / "mesh_log.txt").read_text(errors="ignore"))
items = [("Mesh", f"poly-hexcore, {mesh_rep['cells']:,} cells (clean {int(_cc[-1]):,}), {m['boundary_layers']} prism layers "
                  f"(first cell {m['bl_first_height_mm']} mm), surface {m['min_size_mm']:g}–{m['aircraft_max_size_mm']:g} mm"),
         ("Domain", "box 4 m upstream, 8 m downstream, 4 m to the sides and top/bottom"),
         ("Flow", f"{V:g} m/s, sea-level air (ρ 1.225 kg/m³), steady, incompressible"),
         ("Model", "pressure-based coupled solver, k-ω SST turbulence"),
         ("Angles", f"{', '.join(f'{x:g}' for x in ANGLES)}° by tilting the free stream; one mesh for all"),
         ("Convergence", "stop when CL and CD change < 0.2% per 50 iterations (150–600)")]
y = Inches(1.85)
for k, v in items:
    text(s, Inches(0.6), y, Inches(1.6), Inches(0.7), k, size=15, bold=True, color=POD)
    text(s, Inches(2.25), y, Inches(4.3), Inches(0.75), v, size=14)
    y += Inches(0.85)

# 5 polar charts
s = new_slide(prs, "Lift is almost identical; drag is higher at every angle, so L/D drops.")
title(s, "Aerodynamic coefficients: clean vs pod", f"Whole aircraft, {V:g} m/s")
cw3 = Inches(3.95)
from pptx.enum.chart import XL_LABEL_POSITION as LP
for i, (key, ttl, fmt, lp) in enumerate([("CL", "Lift coefficient CL", "0.00", (LP.ABOVE, LP.BELOW)),
                                         ("CD", "Drag coefficient CD", "0.000", (LP.BELOW, LP.ABOVE)),
                                         ("L_over_D", "Lift-to-drag L/D", "0.0", (LP.ABOVE, LP.BELOW))]):
    line_chart(s, Inches(0.6) + i * (cw3 + Inches(0.13)), Inches(1.75), cw3, Inches(4.4), ttl,
               [("Clean", [g(fc, x, key) for x in ANGLES]), ("Pod", [g(fp, x, key) for x in ANGLES])], fmt, lp)
notes8 = [x for x in ANGLES if "NOT" in fp[x]["status"] or "NOT" in fc[x]["status"]]
foot = ("Unconverged angles (" + ", ".join(f"{x:g}°" for x in notes8) +
        ") report the mean of the last 200 iterations: the flow there is unsteady, likely early separation."
        ) if notes8 else "All angles converged."
text(s, Inches(0.6), Inches(6.45), Inches(12.1), Inches(0.6), foot, size=12, color=MUTED)

# 6 forces table
s = new_slide(prs, "Full table. Forces in newtons at 20 m/s for this aircraft size.")
title(s, "Force results", "Reference area 0.2207 m², chord 0.1641 m, moment centre at the wing quarter chord")
hdr = ("AoA", "Config", "CL", "CD", "Cm", "L/D", "Lift N", "Drag N", "Status")
rows = [hdr]
for x in ANGLES:
    for name, tblf in (("Clean", fc), ("Pod", fp)):
        r = tblf[x]
        rows.append((f"{x:g}°", name, f"{float(r['CL']):.3f}", f"{float(r['CD']):.4f}", f"{float(r['Cm']) + 0.0:.3f}".replace("-0.000", "0.000"),
                     f"{float(r['L_over_D']):.1f}", f"{float(r['lift_N']):.1f}", f"{float(r['drag_N']):.2f}",
                     "converged" if r["status"] == "converged" else "unsteady (mean)"))
tbl = s.shapes.add_table(len(rows), len(hdr), Inches(0.6), Inches(1.8), Inches(12.1), Inches(4.6)).table
widths = [0.9, 1.2, 1.2, 1.3, 1.2, 1.1, 1.3, 1.3, 2.6]
for c, wd in enumerate(widths):
    tbl.columns[c].width = Inches(wd)
for r, row in enumerate(rows):
    for c, val in enumerate(row):
        cell = tbl.cell(r, c)
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(0x3A, 0x3E, 0x43) if r == 0 else PANEL
        cell.text = val
        f = cell.text_frame.paragraphs[0].runs[0].font
        f.size, f.name, f.bold = Pt(14), FONT, r == 0
        f.color.rgb = TEXT if r == 0 or c != 1 else (POD if val == "Pod" else CLN)

# 7 surface pressure across angles
s = new_slide(prs, "Blue = suction (low pressure) on the wing's upper surface, the source of lift; it grows with angle. "
                   "Red = stagnation at the nose, pod front and leading edges.")
title(s, "Surface pressure grows with angle of attack", "Pressure coefficient on the pod configuration")
iw = Inches(3.95)
for i, x in enumerate(ANGLES):
    image(s, render(HERE, x, "2-surface-pressure", "iso"), Inches(0.6) + i * (iw + Inches(0.13)), Inches(1.9),
          w=iw, caption=f"{x:g}°  ·  CL {g(fp, x, 'CL'):.2f}")
image(s, render(HERE, 4, "2-surface-pressure", "top"), Inches(0.6), Inches(4.75), h=Inches(2.3))
image(s, render(HERE, 4, "2-surface-pressure", "bottom"), Inches(4.9), Inches(4.75), h=Inches(2.3))
text(s, Inches(9.35), Inches(4.85), Inches(3.4), Inches(2.2),
     [[("4°: top and underside", {"bold": True})],
      "Suction covers most of the upper wing; the underside sits near free-stream pressure, apart from the pod nose."],
     size=13, color=MUTED)

# 8 the pod
s = new_slide(prs, "The pod's blunt front stagnates the flow (red); its chamfered sides accelerate it (blue), "
                   "and the flat rear leaves a small low-pressure wake. That is where the extra drag comes from.")
title(s, "Flow around the pod", "Underside close-up at 4°")
image(s, render(HERE, 4, "2-surface-pressure", "pod"), Inches(0.6), Inches(1.8), w=Inches(6.0),
      caption="Pressure coefficient")
image(s, render(HERE, 4, "3-streamlines", "pod"), Inches(6.75), Inches(1.8), w=Inches(6.0),
      caption="Streamlines released ahead of the nose")
text(s, Inches(0.6), Inches(5.6), Inches(12.1), Inches(1.3),
     [[("Design hint: ", {"bold": True, "color": POD}),
       ("a rounder nose and a tapered tail on the pod would cut its pressure drag; mounting it closer "
        "to the fuselage reduces the exposed area.", {})]], size=15)

# 9 streamlines
s = new_slide(prs, "Streamlines bend upward ahead of the wing (upwash) and downward behind it (downwash). "
                   "Turning the air downward is what produces lift.")
title(s, "Streamlines over the aircraft", "Colour = local speed, 4°")
image(s, render(HERE, 4, "3-streamlines", "iso"), Inches(0.6), Inches(1.75), w=Inches(8.2))
image(s, render(HERE, 4, "3-streamlines", "top"), Inches(9.0), Inches(1.75), w=Inches(3.75), caption="Top view")
text(s, Inches(9.0), Inches(4.2), Inches(3.75), Inches(2.6),
     ["Flow speeds up over the wing (orange) and slows near the nose and pod (blue).",
      "The tail boom and V-tail sit in the wing's downwash."], size=14, color=MUTED)

# 10 wing section
s = new_slide(prs, "Airfoil section at 0.55 m span. The suction peak strengthens and moves forward as the angle rises.")
title(s, "Wing section: flow speed", "Cut at 0.55 m span, outboard of the motor boom")
def section(path, x, y, w, caption):
    """wing-section render zoomed on the airfoil (crops away the empty far field)"""
    cl, cr, ct, cb = 0.2, 0.12, 0.22, 0.22
    pic = slide_pic = s.shapes.add_picture(str(path), x, y)
    pic.crop_left, pic.crop_right, pic.crop_top, pic.crop_bottom = cl, cr, ct, cb
    pic.width = w
    pic.height = int(w * (9 * (1 - ct - cb)) / (16 * (1 - cl - cr)))
    text(s, x, y + pic.height + Inches(0.08), w, Inches(0.35), caption, size=12, color=MUTED, align=PP_ALIGN.CENTER)
    return pic


big = section(render(HERE, 4, "4-wing-section", "section"), Inches(0.6), Inches(1.75), Inches(7.6),
              "4° · colour = speed, 0 (blue) to 28 m/s (red), free stream 20 m/s")
for i, x in enumerate((ANGLES[0], ANGLES[-1])):
    section(render(HERE, x, "4-wing-section", "section"), Inches(8.85), Inches(1.75) + i * Inches(2.3),
            Inches(3.9), f"{x:g}°")
text(s, Inches(0.6), Inches(6.45), Inches(12.1), Inches(0.8),
     ["Fast flow over the upper surface (red) is the suction that lifts the wing; the thin trailing wake is "
      "profile drag. At 8° the upper boundary layer thickens toward the trailing edge, matching the unsteady forces."],
     size=14)

# 11 vortices & wake
s = new_slide(prs, "Higher pressure under the wing spills around the tips and rolls up into two counter-rotating "
                   "vortices: the source of induced drag.")
title(s, "Wing-tip vortices and wake", "4°, pod configuration")
image(s, render(HERE, 4, "5-tip-vortices", "tip"), Inches(0.6), Inches(1.8), w=Inches(6.0),
      caption="Tip streamlines, coloured by streamwise vorticity")
image(s, render(HERE, 4, "6-wake", "wake"), Inches(6.75), Inches(1.8), w=Inches(6.0),
      caption="Streamwise vorticity 0.1 m and 0.4 m behind the wing")
text(s, Inches(0.6), Inches(5.6), Inches(12.1), Inches(1.3),
     "Red and blue mark opposite rotation. The pair weakens downstream; part of that is the coarse far-field mesh "
     "(student cell limit), so the far wake is under-resolved.", size=15)

# 12 clean vs pod
s = new_slide(prs, "Same angle, same colour scale. The pod is the only obvious change in the surface pressure.")
title(s, "Clean vs pod at 4°", "Underside pressure coefficient, same colour scale")
image(s, render(CLEAN, 4, "2-surface-pressure", "bottom"), Inches(0.6), Inches(1.8), w=Inches(6.0),
      caption=f"Clean · CD {g(fc, 4, 'CD'):.4f} · L/D {g(fc, 4, 'L_over_D'):.1f}")
image(s, render(HERE, 4, "2-surface-pressure", "bottom"), Inches(6.75), Inches(1.8), w=Inches(6.0),
      caption=f"Pod · CD {g(fp, 4, 'CD'):.4f} · L/D {g(fp, 4, 'L_over_D'):.1f}")
text(s, Inches(0.6), Inches(5.75), Inches(12.1), Inches(1.2),
     [[("Everything else looks the same: ", {"bold": True, "color": POD}),
       (f"the wing loading barely changes (CL {g(fc, 4, 'CL'):.3f} vs {g(fp, 4, 'CL'):.3f}), while the pod's "
        f"stagnation front and suction edges plus the boom-front motors add "
        f"{g(fp, 4, 'drag_N') - g(fc, 4, 'drag_N'):.2f} N of drag at {V:g} m/s.", {})]], size=15)

# 13 limits & next steps
s = new_slide(prs, "Close with what the results can and cannot support.")
title(s, "Confidence and next steps")
panel(s, Inches(0.6), Inches(1.6), Inches(5.9), Inches(4.4))
panel(s, Inches(6.85), Inches(1.6), Inches(5.9), Inches(4.4))
text(s, Inches(0.9), Inches(1.85), Inches(5.3), Inches(4.8),
     [[("Good for", {"bold": True, "size": 20, "color": CLN})],
      "Comparing configurations: both use identical mesh settings, solver and reference values.",
      "Trends with angle of attack, pressure distribution, where drag comes from.",
      "Ballpark L/D and lift at cruise."], size=15)
text(s, Inches(7.15), Inches(1.85), Inches(5.3), Inches(4.8),
     [[("Limits", {"bold": True, "size": 20, "color": POD})],
      f"~{mesh_rep['cells'] / 1e6:.2f}M cells (student limit): absolute drag is approximate; the far wake is coarse.",
      "Propellers not modelled (motors as static shapes).",
      "Small gaps (hatch seams, under 24 mm slots) sealed by the wrap.",
      "8°: unsteady, reported as a time average.",
      [("Next: ", {"bold": True}), ("refine near the pod, add a 10–12° point for stall, "
                                     "and try a faired pod shape.", {})]], size=15)

prs.save(OUT)
print(f"wrote {OUT}  ({len(prs.slides)} slides)")
